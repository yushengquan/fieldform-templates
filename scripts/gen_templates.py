#!/usr/bin/env python3
"""从 src/data/trades/*.json 批量生成 xlsx/docx/pdf 模板文件到 public/downloads/。
用法: python3 scripts/gen_templates.py
新增一个工种 = 在 src/data/trades/ 加一个 JSON，再跑本脚本，文件自动生成。"""
import json, datetime, glob, os, re

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA_DIR = os.path.join(BASE, "src", "data", "trades")
OUT_DIR = os.path.join(BASE, "public", "downloads")
os.makedirs(OUT_DIR, exist_ok=True)

SAMPLE_DATE = datetime.date(2026, 9, 29)


def load_trades():
    trades = []
    for path in glob.glob(os.path.join(DATA_DIR, "*.json")):
        with open(path, encoding="utf-8") as f:
            cfg = json.load(f)
        cfg["_path"] = path
        trades.append(cfg)
    return trades


def trade_label(cfg):
    doc = cfg.get("docType", "invoice")
    return f"{cfg.get('title', 'Template').upper()}"


# ---------------- Excel ----------------
def gen_xlsx(cfg, out_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Side, Border
    from openpyxl.utils import get_column_letter

    HEADER = "17365D"; SECOND = "5B9BD5"; LINE = "B8C7D9"
    thin = Side(style="thin", color=LINE)
    top_line = Side(style="thin", color="9AA7B5")
    white_bold = Font(name="Arial", bold=True, color="FFFFFF", size=12)
    title_font = Font(name="Arial", bold=True, color="FFFFFF", size=16)
    label_font = Font(name="Arial", bold=True, size=10, color="17365D")
    small_gray = Font(name="Arial", size=9, color="6B7280")
    base_font = Font(name="Arial", size=10)
    bold_font = Font(name="Arial", size=10, bold=True)
    money_fmt = '"$"#,##0.00'
    date_fmt = "yyyy-mm-dd"

    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEF", (6, 40, 8, 16, 14, 14)):
        ws.column_dimensions[col].width = w

    def fill(addr, value, font=base_font, align=None, numfmt=None, fill_color=None):
        c = ws[addr]
        c.value = value
        c.font = font
        c.alignment = align or Alignment(vertical="center", wrap_text=True)
        if numfmt:
            c.number_format = numfmt
        if fill_color:
            c.fill = PatternFill("solid", fgColor=fill_color)
        return c

    ws.merge_cells("A1:F1")
    fill("A1", trade_label(cfg), title_font, Alignment(horizontal="center", vertical="center"), fill_color=HEADER)
    ws.row_dimensions[1].height = 30

    for addr, text, font in [
        ("A4", "Your Business Name", label_font),
        ("A5", "Street Address, City, State ZIP", base_font),
        ("A6", "Phone  |  Email", base_font),
        ("A7", "License #: ________________", base_font),
        ("A8", "EIN / Tax ID: ______________", base_font),
    ]:
        fill(addr, text, font, Alignment(vertical="center"))
    for r in range(4, 9):
        ws.row_dimensions[r].height = 17

    fill("D4", "INVOICE #:", label_font)
    fill("E4", "INV-0001", base_font)
    fill("D5", "DATE:", label_font)
    fill("E5", SAMPLE_DATE, base_font, Alignment(horizontal="center", vertical="center"), date_fmt)
    fill("D6", "PAYMENT TERMS:", label_font)
    fill("E6", "Net 14", base_font)
    fill("D7", "DUE DATE:", label_font)
    fill("E7", "=E5+14", base_font, Alignment(horizontal="center", vertical="center"), date_fmt)

    fill("A10", "BILL TO", label_font)
    for i, txt in enumerate(["Client Name", "Client Address", "Job Site Address", "Phone  |  Email"]):
        fill(f"B{11+i}", txt, base_font)
        ws.row_dimensions[11+i].height = 16

    hdr_row = 16
    for i, h in enumerate(["#", "Description", "Qty", "Unit", "Rate ($)", "Amount ($)"]):
        col = get_column_letter(i + 1)
        fill(f"{col}{hdr_row}", h, white_bold,
             Alignment(horizontal="center" if i in (0, 2, 3) else "left", vertical="center"),
             fill_color=SECOND)
        ws[f"{col}{hdr_row}"].border = Border(bottom=thin)
    ws.row_dimensions[hdr_row].height = 20

    items = cfg["lineItems"]
    data_start = hdr_row + 1
    for idx, it in enumerate(items):
        r = data_start + idx
        ws.row_dimensions[r].height = 18
        fill(f"A{r}", idx + 1, base_font, Alignment(horizontal="center", vertical="center"))
        fill(f"B{r}", it["description"], base_font)
        fill(f"C{r}", it["defaultQty"], base_font, Alignment(horizontal="center", vertical="center"), "0.##")
        fill(f"D{r}", it["unit"], base_font, Alignment(horizontal="center", vertical="center"))
        fill(f"E{r}", it["sampleRate"], base_font, Alignment(horizontal="right", vertical="center"), money_fmt)
        fill(f"F{r}", f"=C{r}*E{r}", base_font, Alignment(horizontal="right", vertical="center"), money_fmt)
    last = data_start + len(items) - 1

    sub_row = last + 2
    fill(f"D{sub_row}", "SUBTOTAL", label_font)
    fill(f"F{sub_row}", f"=SUM(F{data_start}:F{last})", bold_font,
         Alignment(horizontal="right", vertical="center"), money_fmt)
    ws[f"F{sub_row}"].border = Border(top=top_line)

    tax_rate_row = sub_row + 1
    fill(f"D{tax_rate_row}", "TAX RATE", label_font)
    fill(f"E{tax_rate_row}", None, base_font, Alignment(horizontal="right", vertical="center"), "0.00%")
    fill(f"C{tax_rate_row}", "look up your state rate", small_gray, Alignment(horizontal="left", vertical="center"))
    tax_row = sub_row + 2
    fill(f"D{tax_row}", "TAX", label_font)
    fill(f"F{tax_row}", f'=IF(E{tax_rate_row}="","",ROUND(F{sub_row}*E{tax_rate_row},2))',
         base_font, Alignment(horizontal="right", vertical="center"), money_fmt)
    total_row = sub_row + 3
    fill(f"D{total_row}", "TOTAL DUE", Font(name="Arial", bold=True, size=11, color="17365D"))
    fill(f"F{total_row}", f"=F{sub_row}+IF(ISNUMBER(F{tax_row}),F{tax_row},0)",
         Font(name="Arial", bold=True, size=11, color="17365D"),
         Alignment(horizontal="right", vertical="center"), money_fmt)
    ws[f"F{total_row}"].border = Border(top=top_line, bottom=top_line)

    note_start = total_row + 2
    fill(f"A{note_start}", "PAYMENT TERMS:", label_font)
    fill(f"B{note_start}", cfg.get("paymentTerms", ""), base_font)
    ws.row_dimensions[note_start].height = 30
    fill(f"A{note_start+1}", "NOTES:", label_font)
    fill(f"B{note_start+1}", cfg.get("notes", ""), base_font)

    disc = note_start + 3
    ws.merge_cells(f"A{disc}:F{disc+2}")
    fill(f"A{disc}", cfg.get("disclaimer", ""), small_gray, Alignment(vertical="top", wrap_text=True))

    wb.save(out_path)


# ---------------- Word ----------------
def gen_docx(cfg, out_path):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)

    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run(trade_label(cfg))
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(0x17, 0x36, 0x5D)

    for line in ["Your Business Name", "Street Address, City, State ZIP", "Phone  |  Email",
                 "License #: ________________", "EIN / Tax ID: ______________"]:
        p = doc.add_paragraph(line)
        p.paragraph_format.space_after = Pt(0)

    doc.add_paragraph()
    doc.add_paragraph("Invoice #: INV-0001    |    Date: 2026-09-29    |    Payment Terms: Net 14    |    Due: 2026-10-13")
    bt = doc.add_paragraph()
    bt.add_run("BILL TO").bold = True
    for line in ["Client Name", "Client Address", "Job Site Address", "Phone  |  Email"]:
        doc.add_paragraph(line).paragraph_format.space_after = Pt(0)

    doc.add_paragraph()
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, txt in enumerate(["Description", "Qty", "Unit", "Rate ($)", "Amount ($)"]):
        run = table.rows[0].cells[i].paragraphs[0].add_run(txt)
        run.bold = True

    subtotal = 0.0
    for it in cfg["lineItems"]:
        amt = (it.get("defaultQty", 0)) * (it.get("sampleRate", 0.0))
        subtotal += amt
        cells = table.add_row().cells
        for i, v in enumerate([it["description"], str(it.get("defaultQty", 0)), it["unit"],
                               f'{it.get("sampleRate", 0.0):.2f}', f'{amt:.2f}']):
            cells[i].text = v
    sum_row = table.add_row().cells
    sum_row[0].text = "SUBTOTAL"
    sum_row[4].text = f"${subtotal:.2f}"
    tax_row = table.add_row().cells
    tax_row[0].text = "TAX (enter your state rate)"
    total_row = table.add_row().cells
    total_row[0].text = "TOTAL DUE"
    total_row[4].text = f"${subtotal:.2f}"

    doc.add_paragraph()
    doc.add_paragraph(f"Payment terms: {cfg.get('paymentTerms', '')}")
    doc.add_paragraph(f"Notes: {cfg.get('notes', '')}")
    doc.add_paragraph(cfg.get("disclaimer", ""))
    doc.save(out_path)


# ---------------- PDF ----------------
def gen_pdf(cfg, out_path):
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.colors import HexColor
    from reportlab.pdfgen import canvas

    pdf = canvas.Canvas(out_path, pagesize=letter)
    W, H = letter
    dark = HexColor("#17365D")
    mid = HexColor("#5B9BD5")
    gray = HexColor("#6B7280")

    pdf.setFillColor(dark)
    pdf.setFont("Helvetica-Bold", 18)
    pdf.drawCentredString(W / 2, H - 60, trade_label(cfg))

    y = H - 95
    pdf.setFont("Helvetica", 10)
    for line in ["Your Business Name", "Street Address, City, State ZIP",
                 "Phone  |  Email", "License #: ______________", "EIN / Tax ID: ______________"]:
        pdf.drawString(60, y, line)
        y -= 17

    pdf.setFont("Helvetica-Bold", 10)
    pdf.drawString(400, H - 95, "Invoice #: INV-0001")
    pdf.drawString(400, H - 112, "Date: 2026-09-29")
    pdf.drawString(400, H - 129, "Payment Terms: Net 14")
    pdf.drawString(400, H - 146, "Due Date: 2026-10-13")

    # BILL TO 块：标签放在 From 块下方，客户信息再从标签下方排开，避免与 EIN 行重叠
    pdf.setFont("Helvetica-Bold", 10)
    pdf.drawString(60, y - 8, "BILL TO")
    y2 = y - 25
    for line in ["Client Name", "Client Address", "Job Site Address", "Phone  |  Email"]:
        pdf.drawString(60, y2, line)
        y2 -= 17

    ty = y2 - 18
    pdf.setFillColor(mid)
    pdf.rect(60, ty - 20, W - 120, 20, stroke=0, fill=1)
    pdf.setFillColor(HexColor("#FFFFFF"))
    pdf.setFont("Helvetica-Bold", 10)
    for x, txt in [(70, "Description"), (330, "Qty"), (375, "Unit"), (420, "Rate"), (480, "Amount")]:
        pdf.drawString(x, ty - 13, txt)

    ry = ty - 20
    pdf.setFillColor(HexColor("#111111"))
    pdf.setFont("Helvetica", 10)
    subtotal = 0.0
    for it in cfg["lineItems"]:
        amt = (it.get("defaultQty", 0)) * (it.get("sampleRate", 0.0))
        subtotal += amt
        ry -= 18
        pdf.drawString(70, ry, it["description"])
        pdf.drawString(335, ry, str(it.get("defaultQty", 0)))
        pdf.drawString(380, ry, it["unit"])
        pdf.drawString(425, ry, f"${it.get('sampleRate', 0.0):.2f}")
        pdf.drawString(485, ry, f"${amt:.2f}")

    ry -= 8
    pdf.setStrokeColor(HexColor("#9AA7B5"))
    pdf.line(420, ry, W - 60, ry)
    ry -= 16
    pdf.setFont("Helvetica-Bold", 10)
    pdf.drawString(420, ry, "SUBTOTAL")
    pdf.drawString(485, ry, f"${subtotal:.2f}")
    ry -= 16
    pdf.setFont("Helvetica", 10)
    pdf.drawString(420, ry, "TAX (enter state rate)")
    ry -= 20
    pdf.setFont("Helvetica-Bold", 11)
    pdf.drawString(420, ry, "TOTAL DUE")
    pdf.drawString(485, ry, f"${subtotal:.2f}")

    ry -= 40
    pdf.setFillColor(gray)
    pdf.setFont("Helvetica", 8)
    pdf.drawString(60, ry, f"Payment terms: {cfg.get('paymentTerms', '')}")
    pdf.drawString(60, ry - 12, f"Notes: {cfg.get('notes', '')}")
    pdf.drawString(60, ry - 28, cfg.get("disclaimer", ""))
    pdf.save()


def slugify(name):
    return re.sub(r"[^a-z0-9-]", "", os.path.splitext(os.path.basename(name))[0].lower())


def main():
    trades = load_trades()
    for cfg in trades:
        slug = slugify(cfg["_path"])
        gen_xlsx(cfg, os.path.join(OUT_DIR, f"{slug}.xlsx"))
        gen_docx(cfg, os.path.join(OUT_DIR, f"{slug}.docx"))
        gen_pdf(cfg, os.path.join(OUT_DIR, f"{slug}.pdf"))
        print(f"generated {slug}.xlsx/.docx/.pdf")
    print(f"done: {len(trades)} template(s) -> {OUT_DIR}")


if __name__ == "__main__":
    main()
